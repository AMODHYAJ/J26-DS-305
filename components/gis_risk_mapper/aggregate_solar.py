"""C3 - Aggregate CEB embedded-generation solar capacity per transformer (SIN).

Input : CEB CSC exports (EMB-EMB / EMB-KOL / EMB-GOD / EMB-UDA .xlsx)
Output: one row per transformer with total solar kW, kVA rating, ratio, scheme breakdown.
Cross-checks the SIN-level summary table against the account-level detail table.
"""
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ZW = re.compile(r"[\u200b\u200c\u200d\ufeff]")  # zero-width chars hidden in Section codes

# column indexes (0-based) in the CEB export
SUM_COLS = dict(gss=2, feeder=4, section=8, sin=10, kva=13, scheme=15, kw=20)
DET_COLS = dict(gss=1, feeder=3, section=7, sin=9, account=12, kw=14, date=16, scheme=18)


def clean(v):
    return ZW.sub("", str(v)).strip() if v is not None else None


def sin_num(v):
    """'R011' and '011' -> 11 (detail table has digits only)."""
    m = re.search(r"(\d+)", clean(v) or "")
    return int(m.group(1)) if m else None


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def read_file(path):
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    csc, mode, summ, det = None, None, [], []
    for r in rows:
        first = clean(r[0])
        if first == "CSC :":
            csc = clean(r[5]) or clean(r[3])
        elif first == "No":
            mode = "sum" if clean(r[13]) == "SIN Cap." else "det"
        elif first and first.isdigit() and mode:
            if mode == "sum":
                d = {k: r[i] for k, i in SUM_COLS.items()}
                d["csc"] = csc
                summ.append(d)
            else:
                d = {k: r[i] for k, i in DET_COLS.items()}
                d["csc"] = csc
                det.append(d)
    return pd.DataFrame(summ), pd.DataFrame(det)


def tidy(df):
    for c in ("gss", "feeder", "section", "scheme", "csc"):
        if c in df:
            df[c] = df[c].map(clean)
    df["sin"] = df["sin"].map(clean)          # keep the real letter: R/U/S/W + digits
    df["sin_num"] = df["sin"].map(sin_num)
    df["kw"] = df["kw"].map(num)
    if "kva" in df:
        df["kva"] = df["kva"].map(num)
    return df.dropna(subset=["sin", "kw"])


def main(folder, out_csv):
    s_all, d_all = [], []
    for f in sorted(Path(folder).glob("EMB-*.xlsx")):
        s, d = read_file(f)
        s_all.append(s)
        d_all.append(d)
    S, D = tidy(pd.concat(s_all)), tidy(pd.concat(d_all))
    key = ["csc", "gss", "feeder", "section", "sin_num"]

    # transformer total = sum of all scheme capacities on the same SIN
    tr = S.groupby(key, as_index=False).agg(sin=("sin", "first"), kva=("kva", "max"), solar_kw=("kw", "sum"))
    tr["solar_to_kva"] = (tr["solar_kw"] / tr["kva"]).round(3)
    sch = S.pivot_table(index=key, columns="scheme", values="kw", aggfunc="sum", fill_value=0).reset_index()
    tr = tr.merge(sch, on=key)

    # cross-check against account-level detail
    chk = D.groupby(key, as_index=False).agg(detail_kw=("kw", "sum"), n_accounts=("kw", "size"))
    tr = tr.merge(chk, on=key, how="left")
    tr["match"] = (tr["solar_kw"] - tr["detail_kw"]).abs() < 0.05

    tr = tr[["csc","gss","feeder","section","sin"] + [c for c in tr.columns if c not in ("csc","gss","feeder","section","sin","sin_num")]]
    tr.to_csv(out_csv, index=False)
    print(f"transformers: {len(tr)} | total solar kW: {tr.solar_kw.sum():.1f}")
    print(f"summary vs detail match: {tr.match.sum()}/{len(tr)}")
    bad = tr[~tr.match]
    if len(bad):
        print("MISMATCHES:\n", bad[["csc", "section", "sin", "solar_kw", "detail_kw"]].to_string(index=False))
    # SIN codes repeating across different sections need a section-aware join later
    print(f"unique SIN codes: {tr['sin'].nunique()} of {len(tr)} transformers")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "data/ceb_raw",
         sys.argv[2] if len(sys.argv) > 2 else "solar_by_transformer.csv")